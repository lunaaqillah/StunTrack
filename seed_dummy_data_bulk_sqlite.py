"""Generator data dummy menggunakan SQLite langsung.

Skema dan nama kolom mengikuti model aplikasi agar data hasil seed tetap
kompatibel dengan aplikasi Flask.
"""

# ==========================================
# SECTION: Imports
# ==========================================

import json
import random
import sqlite3
import uuid
from collections import defaultdict
from datetime import date
from dateutil.relativedelta import relativedelta

import openpyxl
from backend.data.variable_database import KUESIONER_ITEMS, ITEM_BY_ID
from backend.routes.survei.scoring import _item_srs_value

with open('backend/data/riwayat_gondok.json', encoding='utf-8') as f:
    RIWAYAT_GONDOK_CONFIG = json.load(f)

with open('backend/data/coefficients.json', encoding='utf-8') as f:
    COEF_BUNDLE = json.load(f)
with open('backend/data/status_risiko.json', encoding='utf-8') as f:
    STATUS_RISIKO_CONFIG = json.load(f)
with open('backend/data/perbandingan_matrix.json', encoding='utf-8') as f:
    PERBANDINGAN_MATRIX = json.load(f)

COEF_MEAN = COEF_BUNDLE['standardization']['mean']
COEF_STD = COEF_BUNDLE['standardization']['std']
COEF_Y = COEF_BUNDLE['model_Y']['coefficients']
COEF_INTERCEPT = COEF_BUNDLE['model_Y']['intercept']


# ==========================================
# SECTION: Configuration
# ==========================================

random.seed(2026)

TODAY = date(2026, 7, 29)
DB_PATH = "backend/StunTrack.db"


PEKERJAAN_OPTIONS = [
    "Pegawai Negeri (PNS/TNI/Polri)",
    "Pegawai Swasta/Karyawan Kantor",
    "Wiraswasta/Pengusaha",
    "Pedagang/Penjual",
    "Petani/Peternak/Perkebun",
    "Nelayan",
    "Buruh/Pekerja Harian/Kuli",
    "Sopir/Ojek/Transportasi",
    "Guru/Dosen/Tenaga Pendidik",
    "Tenaga Kesehatan (Dokter, Perawat, Bidan, dsb.)",
    "Ibu Rumah Tangga",
    "Pensiunan",
    "Tidak Bekerja",
    "Lainnya",
]

PENDIDIKAN_OPTIONS = ["SD", "SMP", "SMA", "Diploma", "Sarjana", "Magister", "Doktor", "Lainnya"]

RIWAYAT_GONDOK_WEIGHTS = list(RIWAYAT_GONDOK_CONFIG['seed_weights'].items())

MALE_FIRST = [
    "Budi", "Agus", "Wahyu", "Dedi", "Bambang", "Eko", "Hendra", "Rudi", "Joko", "Slamet",
    "Rizky", "Fajar", "Andi", "Dwi", "Yusuf", "Fadli", "Irfan", "Gilang", "Rio", "Arif",
    "Taufik", "Hadi", "Sutrisno", "Herman", "Sugianto", "Anwar", "Nanang", "Ridho", "Yoga", "Bayu",
]
FEMALE_FIRST = [
    "Siti", "Sri", "Dewi", "Rina", "Wati", "Ani", "Fitri", "Yuni", "Lestari", "Indah",
    "Ratna", "Nur", "Putri", "Wulan", "Diah", "Rahayu", "Kartika", "Mega", "Yanti", "Sari",
    "Nia", "Lina", "Eka", "Novi", "Dian", "Retno", "Suryani", "Handayani", "Wahyuni", "Puspita",
]
SURNAMES = [
    "Santoso", "Wijaya", "Kusuma", "Saputra", "Pratama", "Setiawan", "Susanto", "Hidayat",
    "Nugroho", "Rahayu", "Gunawan", "Firmansyah", "Permana", "Wibowo", "Purnomo", "Halim",
    "Simanjuntak", "Siregar", "Napitupulu", "Panjaitan", "Manurung", "Situmorang",
    "Wulandari", "Anggraini", "Maharani", "Ramadhan", "Iskandar", "Prasetyo", "Sulistiawati", "Handoko",
]
KOTA = [
    "Malang", "Surabaya", "Blitar", "Kediri", "Jember", "Sidoarjo", "Gresik", "Mojokerto",
    "Pasuruan", "Probolinggo", "Banyuwangi", "Tulungagung", "Madiun", "Bojonegoro", "Lamongan",
    "Yogyakarta", "Semarang", "Solo", "Bandung", "Cirebon", "Bekasi", "Bogor", "Tangerang",
    "Denpasar", "Makassar", "Medan", "Palembang", "Balikpapan", "Pontianak", "Manado",
]
JALAN = [
    "Jl. Mawar", "Jl. Melati", "Jl. Anggrek", "Jl. Kenanga", "Jl. Flamboyan", "Jl. Kartini",
    "Jl. Diponegoro", "Jl. Sudirman", "Jl. Ahmad Yani", "Jl. Gajah Mada", "Jl. Veteran",
    "Jl. Merdeka", "Jl. Pahlawan", "Jl. Cempaka", "Jl. Dahlia", "Jl. Teratai",
]

CHILD_MALE_FIRST = [
    "Rangga", "Bima", "Arka", "Fathan", "Zidan", "Aditya", "Raka", "Farel", "Gading", "Kenzo",
    "Aksa", "Danish", "Rasyid", "Aril", "Fikri",
]
CHILD_FEMALE_FIRST = [
    "Kirana", "Alya", "Zahra", "Chelsea", "Naura", "Cinta", "Keysha", "Aqila", "Bunga", "Ayu",
    "Salsa", "Nadin", "Marsha", "Aurel", "Icha",
]

SKALA_WEIGHTS = {
    "Normal":         [2, 3, 4, 5, 6],
    "Pendek":         [3, 4, 4, 4, 3],
    "Sangat Pendek":  [6, 5, 4, 3, 2],
}
SKALA_WEIGHTS_X10 = {
    "Normal":         [2, 2, 3, 5, 7],
    "Pendek":         [3, 4, 5, 3, 2],
    "Sangat Pendek":  [7, 5, 3, 2, 2],
}


# ==========================================
# SECTION: Helper Functions
# ==========================================
# ==================== 1. MEMVARIASIKAN WEIGHT DATA ====================

def jitter_weights(weights, spread):
    lo, hi = spread
    return [max(0.05, w * random.uniform(lo, hi)) for w in weights]


STORYLINES = [
    (["Normal"], 0.38),
    (["Normal", "Normal", "Pendek"], 0.10),
    (["Pendek"], 0.13),
    (["Sangat Pendek"], 0.09),
    (["Normal", "Pendek", "Sangat Pendek"], 0.08),
    (["Sangat Pendek", "Pendek", "Normal"], 0.07),
    (["Pendek", "Pendek", "Normal"], 0.07),
    (["Normal", "Sangat Pendek"], 0.08),
]

TARGET_Z_RANGE = {
    "Normal": (-1.3, 0.6),
    "Pendek": (-2.9, -2.0),
    "Sangat Pendek": (-4.5, -3.05),
}


_wb = openpyxl.load_workbook("backend/data/who_standard_zscore.xlsx")
tbu_boys, tbu_girls = {}, {}
for _row in _wb["TBU_BOYS"].iter_rows(min_row=2, values_only=True):
    tbu_boys[int(_row[0])] = {"L": _row[1], "M": _row[2], "S": _row[3]}
for _row in _wb["TBU_GIRLS"].iter_rows(min_row=2, values_only=True):
    tbu_girls[int(_row[0])] = {"L": _row[1], "M": _row[2], "S": _row[3]}


# ==================== 2. MENGHITUNG Z-SCORE TB MENURUT UMUR ====================

def calculate_zscore_tb(tb, month, gender):
    data = (tbu_boys if gender == "Laki-laki" else tbu_girls).get(month)
    if data is None:
        raise ValueError(f"No WHO data for month {month}")
    L, M, S = data["L"], data["M"], data["S"]
    z = (tb / M - 1) / S if L == 0 else (((tb / M) ** L) - 1) / (L * S)
    return round(z, 3)


# ==================== 3. MENCARI TB UNTUK TARGET Z-SCORE ====================

def find_tb_for_target_zscore(target_z, month, gender, tolerance=0.01, max_iter=100):
    lo, hi = 20.0, 150.0
    for _ in range(max_iter):
        mid = (lo + hi) / 2
        z = calculate_zscore_tb(mid, month, gender)
        if abs(z - target_z) < tolerance:
            return round(mid, 1)
        if z < target_z:
            lo = mid
        else:
            hi = mid
    return round((lo + hi) / 2, 1)


# ==================== 4. MENENTUKAN KLASIFIKASI STUNTING ====================

def classify_stunting(zscore):
    if zscore < -3:
        return "Sangat Pendek"
    elif zscore < -2:
        return "Pendek"
    return "Normal"


# ==================== 5. MEMPERKIRAKAN BERAT BADAN ====================

def approx_bb(age_months, status):
    if age_months <= 12:
        base = 3.3 + age_months * 0.6
    else:
        base = 10.5 + (age_months - 12) * 0.22
    base *= random.uniform(0.94, 1.06)
    if status == "Pendek":
        base *= 0.96
    elif status == "Sangat Pendek":
        base *= 0.92
    return round(max(base, 0.5), 1)


# ==================== 6. MEMILIH DATA BERDASARKAN BOBOT ====================

def weighted_choice(pairs):
    items, weights = zip(*pairs)
    return random.choices(items, weights=weights, k=1)[0]


# ==================== 7. MEMILIH STORYLINE DATA ====================

def pick_storyline():
    items = [s for s, _ in STORYLINES]
    weights = [w for _, w in STORYLINES]
    return random.choices(items, weights=weights, k=1)[0]


# ==================== 8. MEMBENTUK JADWAL USIA PENGISIAN ====================

def build_age_schedule(birth_date, today, max_fills, step_choices=None):
    steps = step_choices or [1, 1, 2, 2, 3, 4, 6]
    ages = []
    age = random.randint(1, 3 if step_choices else 6)
    while age <= 60 and len(ages) < max_fills:
        d = birth_date + relativedelta(months=age)
        if d > today:
            break
        ages.append(age)
        age += random.choice(steps)
    return ages


# ==================== 9. MEMBUAT ALAMAT ACAK ====================

def random_alamat():
    return f"{random.choice(JALAN)} No. {random.randint(1, 200)}, {random.choice(KOTA)}"


# ==================== 10. MEMBUAT DETAIL JAWABAN KUESIONER ====================

def generate_kuesioner_details(kuesioner_items, status, extra_variance=False):
    spread = (0.4, 1.8) if extra_variance else (0.65, 1.4)
    details = []
    for item in kuesioner_items:
        id_item = item["id_item"]
        id_variabel = item["id_variabel"]
        base_weights = (SKALA_WEIGHTS_X10 if id_variabel == "X10" else SKALA_WEIGHTS)[status]
        weights = jitter_weights(base_weights, spread)
        skala = random.choices([1, 2, 3, 4, 5], weights=weights, k=1)[0]
        details.append({"id_item": id_item, "id_skala": skala, "id_variabel": id_variabel,
                        "id_indikator": item["id_indikator"]})
    return details


# ==================== 11. MENGHITUNG RATA-RATA VARIABEL ====================

def compute_variable_averages(details):
    values = {f"X{i}": [] for i in range(1, 11)}
    for d in details:
        var = d["id_variabel"]
        if var in values:
            values[var].append(_item_srs_value(d["id_item"], d["id_skala"]))
    return {var: (sum(vals) / len(vals) if vals else 0.0) for var, vals in values.items()}


# ==================== 12. MENGAMBIL ATAU MEMBUAT DATA ORANG TUA ====================

def get_or_create_orang_tua(cur, nama, tanggal_lahir, pekerjaan, alamat, pendidikan, email, riwayat_gondok):
    cur.execute(
        "SELECT id_orang_tua FROM identitas_orang_tua WHERE email=? AND tanggal_lahir=?",
        (email, tanggal_lahir.isoformat()),
    )
    row = cur.fetchone()
    if row:
        return row[0], False
    id_orang_tua = f"ORTU-{uuid.uuid4().hex}"
    cur.execute(
        """INSERT INTO identitas_orang_tua
           (id_orang_tua, nama_orang_tua, tanggal_lahir, pekerjaan, alamat, pendidikan, email, riwayat_gondok)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        (id_orang_tua, nama, tanggal_lahir.isoformat(), pekerjaan, alamat, pendidikan, email, riwayat_gondok),
    )
    return id_orang_tua, True


# ==================== 13. MENGAMBIL ATAU MEMBUAT DATA BALITA ====================

def get_or_create_balita(cur, id_orang_tua, nama, tanggal_lahir, jenis_kelamin):
    cur.execute(
        "SELECT id_balita FROM identitas_balita WHERE nama_balita=? AND tanggal_lahir=? AND id_orang_tua=?",
        (nama, tanggal_lahir.isoformat(), id_orang_tua),
    )
    row = cur.fetchone()
    if row:
        return row[0], False
    id_balita = f"BLT-{uuid.uuid4().hex}"
    cur.execute(
        """INSERT INTO identitas_balita (id_balita, nama_balita, tanggal_lahir, jenis_kelamin, id_orang_tua)
           VALUES (?, ?, ?, ?, ?)""",
        (id_balita, nama, tanggal_lahir.isoformat(), jenis_kelamin, id_orang_tua),
    )
    return id_balita, True


# ==================== 14. MEMBUAT DATA PENGISIAN BALITA ====================

def create_pengisian_for_balita(cur, id_balita, jenis_kelamin, tanggal_lahir_balita, kuesioner_items,
                                 storyline, min_fills=None, extra_variance=False):
    cur.execute("SELECT id_pengisian FROM pengisian WHERE id_balita=?", (id_balita,))
    old_ids = [r[0] for r in cur.fetchall()]
    if old_ids:
        cur.executemany("DELETE FROM detail_faktor_risiko WHERE id_pengisian=?", [(i,) for i in old_ids])
        cur.executemany("DELETE FROM pengisian WHERE id_pengisian=?", [(i,) for i in old_ids])

    if min_fills:
        max_fills = min_fills
        step_choices = [1, 1, 1, 2]
        ages = build_age_schedule(tanggal_lahir_balita, TODAY, max_fills, step_choices=step_choices)
    else:
        max_fills = random.choice([1, 2, 2, 3, 3, 3, 4, 4, 5, 6, 7, 8])
        ages = build_age_schedule(tanggal_lahir_balita, TODAY, max_fills)
    if not ages:
        return 0, 0

    n_stages = len(storyline)
    prev_tb, prev_bb = 0.0, 0.0
    n_pengisian = 0
    n_detail = 0

    for i, age in enumerate(ages):
        stage_status = storyline[min(i * n_stages // len(ages), n_stages - 1)]
        lo, hi = TARGET_Z_RANGE[stage_status]
        target_z = random.uniform(lo, hi)

        tb = find_tb_for_target_zscore(target_z, age, jenis_kelamin)
        if tb < prev_tb:
            tb = round(prev_tb + 0.1, 1)
        tb = max(40.0, min(130.0, tb))
        prev_tb = tb

        actual_z = calculate_zscore_tb(tb, age, jenis_kelamin)
        actual_status = classify_stunting(actual_z)

        bb = approx_bb(age, actual_status)
        if bb < prev_bb:
            bb = round(prev_bb + 0.1, 1)
        prev_bb = bb

        details = generate_kuesioner_details(kuesioner_items, actual_status, extra_variance=extra_variance)
        vars_avg = compute_variable_averages(details)

        tanggal = tanggal_lahir_balita + relativedelta(months=age)
        id_pengisian = f"ANS-{uuid.uuid4().hex}"

        cur.execute(
            """INSERT INTO pengisian
               (id_pengisian, id_balita, tanggal_pengisian, usia_bulan, tb_balita, bb_balita,
                nilai_tbu, status_pertumbuhan)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                id_pengisian, id_balita, tanggal.isoformat(), age, tb, bb, actual_z, actual_status,
            ),
        )
        n_pengisian += 1

        factor_rows = [
            (id_pengisian, var, vars_avg[var])
            for var in vars_avg
        ]
        cur.executemany(
            """INSERT INTO faktor_risiko (id_pengisian, id_variabel, skor_komposit)
               VALUES (?, ?, ?)""",
            factor_rows,
        )

        detail_rows = [
            (id_pengisian, d["id_indikator"], d["id_item"], d["id_skala"],
             _item_srs_value(d["id_item"], d["id_skala"]))
            for d in details
        ]
        cur.executemany(
            """INSERT INTO detail_faktor_risiko
               (id_pengisian, id_indikator, id_item, id_skala, skor_srs)
               VALUES (?, ?, ?, ?, ?)""",
            detail_rows,
        )

        # Simpan snapshot hasil model agar halaman analisis tidak menghitung ulang.
        contributions, y_pred, status_risiko, perbandingan, status_perbandingan = calculate_prediction_snapshot(vars_avg, actual_status)
        cur.execute(
            "UPDATE pengisian SET nilai_prediksi=?, status_risiko=?, perbandingan=?, status_perbandingan=? WHERE id_pengisian=?",
            (y_pred, status_risiko, int(perbandingan) if perbandingan is not None else None, status_perbandingan, id_pengisian),
        )
        cur.executemany(
            "UPDATE faktor_risiko SET skor_zscore=?, kontribusi=? WHERE id_pengisian=? AND id_variabel=?",
            [(z, kontribusi, id_pengisian, var) for var, z, kontribusi in contributions],
        )
        n_detail += len(detail_rows)

    return n_pengisian, n_detail



def calculate_prediction_snapshot(vars_avg, status_pertumbuhan):
    contributions = []
    for var, koef in COEF_Y.items():
        score = vars_avg.get(var, 0.0)
        mean = COEF_MEAN.get(var)
        std = COEF_STD.get(var)
        if mean is None or not std:
            continue
        z = (score - mean) / std
        contributions.append((var, round(float(z), 3), round(float(koef * z), 4)))

    y_pred = float(COEF_INTERCEPT) + sum(row[2] for row in contributions)
    kategori = None
    for row in STATUS_RISIKO_CONFIG['statuses']:
        lower_ok = row.get('min') is None or (y_pred >= row['min'] if row.get('min_inclusive', False) else y_pred > row['min'])
        upper_ok = row.get('max') is None or (y_pred <= row['max'] if row.get('max_inclusive', False) else y_pred < row['max'])
        if lower_ok and upper_ok:
            kategori = row
            break
    if kategori is None:
        raise ValueError(f'Prediksi risiko tidak dapat diklasifikasikan: {y_pred}')

    matrix_info = PERBANDINGAN_MATRIX.get(status_pertumbuhan, {}).get(kategori['level'])
    return contributions, round(y_pred, 3), kategori['value'], bool(matrix_info['is_selaras']) if matrix_info else None, matrix_info['level'] if matrix_info else None

# ==================== 15. MENJALANKAN PROSES SEED DATA ====================

def main():
    con = sqlite3.connect(DB_PATH)
    cur = con.cursor()

    print("=" * 70)
    print("BULK SEEDING STUNTRACK - 500 ORANG TUA / ~1000 BALITA (raw sqlite3)")
    print("=" * 70)

    kuesioner_items = KUESIONER_ITEMS
    if not kuesioner_items:
        raise SystemExit("Referensi kuesioner kosong.")
    print(f"Kuesioner items loaded: {len(kuesioner_items)}")

    total_parents = 0
    total_balita = 0
    total_pengisian = 0
    total_detail = 0


    print("\n" + "-" * 70)
    print("KELUARGA UTAMA: Intan Putri Setyowati")
    print("-" * 70)

    main_id_ortu, created = get_or_create_orang_tua(
        cur,
        nama="Intan Putri Setyowati",
        tanggal_lahir=date(2005, 7, 1),
        pekerjaan="Ibu Rumah Tangga",
        alamat="Jl. Merdeka No. 123, Malang",
        pendidikan="SMA",
        email="intanptrsy@gmail.com",
        riwayat_gondok="Tidak",
    )
    con.commit()
    print(f"Orang tua utama: {main_id_ortu} ({'baru' if created else 'sudah ada'})")
    total_parents += 1

    main_children = [
        ("Adam Setyawan", date(2020, 1, 1), "Laki-laki",
         ["Normal", "Normal", "Pendek", "Pendek"]),
        ("Bunga Setyawan", date(2021, 2, 2), "Perempuan",
         ["Normal"]),
        ("Cahaya Setyawan", date(2022, 3, 3), "Laki-laki",
         ["Sangat Pendek"]),
        ("Dewi Setyawan", date(2023, 4, 4), "Perempuan",
         ["Pendek", "Normal", "Pendek"]),
        ("Eka Setyawan", date(2024, 5, 5), "Laki-laki",
         ["Normal"]),
    ]

    for nama, tgl_lahir, gender, storyline in main_children:
        id_balita, created = get_or_create_balita(cur, main_id_ortu, nama, tgl_lahir, gender)
        total_balita += 1
        n_p, n_d = create_pengisian_for_balita(
            cur, id_balita, gender, tgl_lahir, kuesioner_items, storyline,
            min_fills=random.randint(12, 18), extra_variance=True,
        )
        total_pengisian += n_p
        total_detail += n_d
        con.commit()
        print(f"  {nama} ({gender}, lahir {tgl_lahir}) -> {n_p} pengisian, {n_d} detail jawaban")


    N_EXTRA_PARENTS = 499
    TARGET_EXTRA_CHILDREN = 995

    big_family_flags = [False] * N_EXTRA_PARENTS
    n_big = max(1, round(N_EXTRA_PARENTS * 0.05))
    for idx in random.sample(range(N_EXTRA_PARENTS), n_big):
        big_family_flags[idx] = True

    family_sizes = [1] * N_EXTRA_PARENTS
    capacity = [(5 if big else 3) for big in big_family_flags]
    remaining_extra = TARGET_EXTRA_CHILDREN - N_EXTRA_PARENTS

    for i in range(N_EXTRA_PARENTS):
        if big_family_flags[i]:
            force_extra = random.choice([4, 5])
            family_sizes[i] += force_extra
            capacity[i] -= force_extra
            remaining_extra -= force_extra

    eligible = [i for i in range(N_EXTRA_PARENTS) if capacity[i] > 0]
    while remaining_extra > 0 and eligible:
        i = random.choice(eligible)
        family_sizes[i] += 1
        capacity[i] -= 1
        remaining_extra -= 1
        if capacity[i] <= 0:
            eligible.remove(i)
    while remaining_extra < 0:
        candidates = [i for i in range(N_EXTRA_PARENTS) if family_sizes[i] > 1 and not big_family_flags[i]]
        if not candidates:
            break
        i = random.choice(candidates)
        family_sizes[i] -= 1
        remaining_extra += 1

    print("\n" + "-" * 70)
    print(f"MEMBUAT {N_EXTRA_PARENTS} ORANG TUA ACAK LAINNYA "
          f"(total anak direncanakan: {sum(family_sizes)})")
    print("-" * 70)

    used_emails = {"intanptrsy@gmail.com"}
    used_child_names = set()

    for p_idx in range(N_EXTRA_PARENTS):
        is_big = big_family_flags[p_idx]
        n_children = family_sizes[p_idx]

        gender_ortu = random.choice(["L", "P"])
        first = random.choice(MALE_FIRST if gender_ortu == "L" else FEMALE_FIRST)
        surname = random.choice(SURNAMES)
        nama_ortu = None if random.random() < 0.05 else f"{first} {surname}"

        if is_big:
            oldest_birth_year = random.randint(2016, 2020)
        else:
            oldest_birth_year = random.randint(2017, 2026)
        oldest_birth_month = random.randint(1, 12)
        oldest_birth_day = random.randint(1, 28)
        try:
            oldest_birth_date = date(oldest_birth_year, oldest_birth_month, oldest_birth_day)
        except ValueError:
            oldest_birth_date = date(oldest_birth_year, oldest_birth_month, 28)
        if oldest_birth_date > TODAY:
            oldest_birth_date = TODAY - relativedelta(months=random.randint(1, 12))

        parent_age_gap = random.randint(18, 34)
        parent_birth_year = min(2005, oldest_birth_date.year - parent_age_gap)
        parent_birth_year = max(parent_birth_year, 1965)
        parent_dob = date(parent_birth_year, random.randint(1, 12), random.randint(1, 28))

        while True:
            email = f"{first.lower()}{surname.lower()}{random.randint(1, 9999)}@{'gmail.com' if random.random() < 0.7 else 'yahoo.com'}"
            if email not in used_emails:
                used_emails.add(email)
                break

        id_ortu, _created = get_or_create_orang_tua(
            cur,
            nama=nama_ortu,
            tanggal_lahir=parent_dob,
            pekerjaan=random.choice(PEKERJAAN_OPTIONS),
            alamat=random_alamat(),
            pendidikan=random.choice(PENDIDIKAN_OPTIONS),
            email=email,
            riwayat_gondok=weighted_choice(RIWAYAT_GONDOK_WEIGHTS),
        )
        total_parents += 1

        child_birthdates = [oldest_birth_date]
        cursor_date = oldest_birth_date
        for _ in range(n_children - 1):
            gap_years = random.uniform(1.2, 3.0)
            cursor_date = cursor_date + relativedelta(months=int(gap_years * 12))
            if cursor_date > TODAY - relativedelta(months=1):
                cursor_date = TODAY - relativedelta(months=random.randint(1, 6))
            child_birthdates.append(cursor_date)

        for c_idx, c_dob in enumerate(child_birthdates):
            gender = random.choice(["Laki-laki", "Perempuan"])
            first_name_pool = CHILD_MALE_FIRST if gender == "Laki-laki" else CHILD_FEMALE_FIRST
            for _try in range(10):
                child_name = f"{random.choice(first_name_pool)} {surname}"
                key = (child_name, c_dob, id_ortu)
                if key not in used_child_names:
                    used_child_names.add(key)
                    break

            id_balita, _created = get_or_create_balita(cur, id_ortu, child_name, c_dob, gender)
            total_balita += 1

            storyline = pick_storyline()
            n_p, n_d = create_pengisian_for_balita(cur, id_balita, gender, c_dob, kuesioner_items, storyline)
            total_pengisian += n_p
            total_detail += n_d

        con.commit()

        if (p_idx + 1) % 50 == 0:
            print(f"  ... {p_idx + 1}/{N_EXTRA_PARENTS} orang tua selesai "
                  f"(balita: {total_balita}, pengisian: {total_pengisian})")

    print("\n" + "=" * 70)
    print("BULK SEEDING SELESAI")
    print("=" * 70)
    print(f"Total orang tua   : {total_parents}")
    print(f"Total balita      : {total_balita}")
    print(f"Total pengisian   : {total_pengisian}")
    print(f"Total detail jwbn : {total_detail}")
    print("\nKeluarga utama tetap bisa dicek lewat email: intanptrsy@gmail.com")

    con.close()


if __name__ == "__main__":
    main()
