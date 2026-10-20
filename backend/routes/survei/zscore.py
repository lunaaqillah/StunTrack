"""Fungsi perhitungan usia dan z-score tinggi badan menurut umur (WHO TB/U).

Referensi WHO dimuat satu kali pada level modul dan digunakan bersama oleh
fitur pengisian, prediksi, dan dashboard.
"""


# ==========================================
# SECTION: Imports
# ==========================================

from datetime import datetime
from pathlib import Path
import json

import openpyxl
from dateutil.relativedelta import relativedelta


# ==========================================
# SECTION: Configuration
# ==========================================

DATA_DIR = Path(__file__).resolve().parents[2] / 'data'
with open(DATA_DIR / 'status_pertumbuhan.json', encoding='utf-8') as f:
    STATUS_PERTUMBUHAN_CONFIG = json.load(f)
STATUS_PERTUMBUHAN_VALUES = [row['value'] for row in STATUS_PERTUMBUHAN_CONFIG['statuses']]

growth_workbook = openpyxl.load_workbook(DATA_DIR / 'who_standard_zscore.xlsx', read_only=True, data_only=True)

tbu_boys = {}
tbu_girls = {}

_ws_tbu_boys = growth_workbook['TBU_BOYS']
for _row in _ws_tbu_boys.iter_rows(min_row=2, values_only=True):
    _month = int(_row[0])
    tbu_boys[_month] = {'L': _row[1], 'M': _row[2], 'S': _row[3]}

_ws_tbu_girls = growth_workbook['TBU_GIRLS']
for _row in _ws_tbu_girls.iter_rows(min_row=2, values_only=True):
    _month = int(_row[0])
    tbu_girls[_month] = {'L': _row[1], 'M': _row[2], 'S': _row[3]}


# ==========================================
# SECTION: Business Logic
# ==========================================
# ==================== 1. MENGHITUNG USIA DALAM BULAN ====================

def calculate_age_in_months(birth_date, current_date=None):
    if current_date is None:
        current_date = datetime.now().date()
    delta = relativedelta(current_date, birth_date)
    age_months = delta.years * 12 + delta.months
    return min(age_months, 60)


# ==================== 2. MENGHITUNG Z-SCORE TB MENURUT UMUR ====================

def calculate_zscore_tb(tb, month, gender):
    if gender == 'Laki-laki':
        data = tbu_boys.get(month)
    elif gender == 'Perempuan':
        data = tbu_girls.get(month)
    else:
        raise ValueError("Invalid gender")
    if data is None:
        raise ValueError(f"No data for month {month}")

    L = data['L']
    M = data['M']
    S = data['S']


    if L == 0:
        z = (tb / M - 1) / S
    else:
        z = (((tb / M) ** L) - 1) / (L * S)
    return round(z, 3)


# ==================== 3. MENENTUKAN KLASIFIKASI STUNTING ====================

def classify_stunting(zscore):
    for rule in STATUS_PERTUMBUHAN_CONFIG['zscore_rules']:
        lower_ok = rule.get('min') is None or (zscore >= rule['min'] if rule.get('min_inclusive', False) else zscore > rule['min'])
        upper_ok = rule.get('max') is None or (zscore <= rule['max'] if rule.get('max_inclusive', False) else zscore < rule['max'])
        if lower_ok and upper_ok:
            return rule['value']
    raise ValueError(f'Z-score pertumbuhan tidak dapat diklasifikasikan: {zscore}')
